#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
تبدیل فایل اکسل محتوا به CSV استاندارد، برای خوراک‌دادن به scripts/import-content.mjs

    prompts_index_with_titles_2.xlsx  →  prompts_index_with_titles_2.csv

چرا این فایل جداست؟
    اسکریپت import عمداً هیچ وابستگی npm ندارد (پارسر CSV دستی نوشته شده)، پس
    xlsx نمی‌خواند. بنابراین اکسل یک‌بار به CSV تبدیل می‌شود و از آن به بعد
    ورودیِ import همان CSV است.

نیازمندی:
    python3 و پکیج openpyxl  →  pip install openpyxl

جایگزین بدون پایتون:
    در خود اکسل: File → Save As → «CSV UTF-8 (Comma delimited) (*.csv)»
    خروجی اکسل هم RFC-4180 است و پارسرِ اسکریپت import آن را می‌فهمد.

اجرا (از ریشه‌ی پروژه):
    python scripts/xlsx-to-csv.py
    python scripts/xlsx-to-csv.py --in my.xlsx --out my.csv --sheet "پرامپت‌ها"

چند نکته‌ی مهم در خروجی:
    * UTF-8 بدون BOM — پارسر import هر دو حالت را تحمل می‌کند.
    * پایان خط CRLF، مطابق RFC-4180.
    * quoting فقط در صورت نیاز. این حیاتی است: در این فایل ۴۴۶ پرامپت داخل
      متن‌شان newline دارند، ۶۸۵ کاما و ۸۶ گیومه. بدون quoting درست، CSV
      کاملاً به‌هم می‌ریزد.
    * نام ستون‌ها عیناً از اکسل کپی می‌شود (از جمله «تاریخ»)، چون import
      ستون‌ها را با نام پیدا می‌کند نه با شماره.
"""

import argparse
import csv
import sys
from pathlib import Path

DEFAULT_SHEET = "پرامپت‌ها"


def cell_to_text(value):
    """هر سلول را به رشته‌ی متنی تبدیل می‌کند."""
    if value is None:
        return ""
    # اگر اکسل ستون تاریخ را به‌صورت datetime تحویل داد، با همان قالبِ
    # dd.mm.yyyy HH:MM:SS می‌نویسیم که import انتظارش را دارد.
    if hasattr(value, "strftime"):
        return value.strftime("%d.%m.%Y %H:%M:%S")
    return str(value)


def main():
    root = Path(__file__).resolve().parent.parent

    ap = argparse.ArgumentParser(add_help=True)
    ap.add_argument("--in", dest="src", default="prompts_index_with_titles_2.xlsx")
    ap.add_argument("--out", dest="dst", default="prompts_index_with_titles_2.csv")
    ap.add_argument("--sheet", dest="sheet", default=DEFAULT_SHEET)
    args = ap.parse_args()

    src = Path(args.src)
    dst = Path(args.dst)
    if not src.is_absolute():
        src = root / src
    if not dst.is_absolute():
        dst = root / dst

    if not src.exists():
        sys.exit(f"✗ فایل اکسل پیدا نشد: {src}")

    try:
        import openpyxl
    except ImportError:
        sys.exit("✗ پکیج openpyxl نصب نیست.  →  pip install openpyxl")

    wb = openpyxl.load_workbook(src, read_only=True, data_only=True)
    if args.sheet not in wb.sheetnames:
        sys.exit(f"✗ شیت «{args.sheet}» نیست. شیت‌های موجود: {wb.sheetnames}")

    rows = list(wb[args.sheet].iter_rows(values_only=True))
    if not rows:
        sys.exit("✗ شیت خالی است.")

    header, data = rows[0], rows[1:]
    # ردیف‌های کاملاً خالیِ انتهای شیت را دور می‌ریزیم.
    data = [r for r in data if any(c is not None and str(c).strip() != "" for c in r)]

    with open(dst, "w", encoding="utf-8", newline="") as f:
        w = csv.writer(f, quoting=csv.QUOTE_MINIMAL, lineterminator="\r\n")
        w.writerow([cell_to_text(c) for c in header])
        for r in data:
            w.writerow([cell_to_text(c) for c in r])

    print(f"ستون‌ها : {' | '.join(cell_to_text(c) for c in header)}")
    print(f"ردیف‌ها : {len(data)}")
    print(f"نوشته شد: {dst}")
    print("\nقدم بعدی:  node scripts/import-content.mjs --dry-run")


if __name__ == "__main__":
    main()
